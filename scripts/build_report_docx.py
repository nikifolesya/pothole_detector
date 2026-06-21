from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "practice_report.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(9.5)
            if row_idx == 0:
                set_cell_shading(cell, "F2F4F7")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.allow_autofit = False
    for idx, text in enumerate(headers):
        table.rows[0].cells[idx].text = text
    for row_data in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row_data):
            cells[idx].text = str(text)
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
    style_table(table)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Отчет по практике")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor.from_string("0B2545")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Сравнение архитектур для детекции дефектов дорожного покрытия").italic = True

    doc.add_heading("Тема и цель", level=1)
    doc.add_paragraph(
        "Тема работы: сравнение архитектур для детекции дефектов дорожного покрытия на изображениях и видео."
    )
    doc.add_paragraph(
        "Цель практики: разработать воспроизводимый пайплайн компьютерного зрения, сравнить несколько архитектур детекции, выбрать пригодное прикладное решение и оформить его как сервис инференса."
    )

    doc.add_heading("Объект, предмет и данные", level=1)
    doc.add_paragraph("Объект исследования: изображения дорожного покрытия с дефектами.")
    doc.add_paragraph(
        "Предмет исследования: методы детекции объектов для локализации дорожных ям, трещин, участков повреждения, люков и визуально похожих классов."
    )
    doc.add_paragraph(
        "Входные данные: RGB-изображение или кадр видео. Выходные данные: класс объекта, уверенность, координаты bounding box, аннотированное изображение и запись в историю запусков."
    )
    add_table(
        doc,
        ["split", "images", "labels", "objects", "empty labels"],
        [
            ["train", "5325", "5325", "13022", "13"],
            ["valid", "4943", "4943", "12042", "18"],
            ["test", "4070", "4070", "10329", "17"],
        ],
        [1.1, 1.0, 1.0, 1.0, 1.2],
    )
    doc.add_paragraph(
        "Датасет Roboflow pothole-detection-2-zvaru version 6 содержит 16 классов. Основные прикладные классы: Pothole, medium-pothole, risk-pothole, pothole_water, crack, damage."
    )

    doc.add_heading("Гипотеза", level=1)
    doc.add_paragraph(
        "Легкие YOLO-модели должны дать лучший баланс качества и скорости для сервиса, а RT-DETR может повысить качество на сложных сценах ценой большего времени инференса. Увеличение входного разрешения должно улучшить обнаружение малых ям, но ухудшить latency."
    )

    doc.add_heading("Пайплайн работ", level=1)
    add_bullets(
        doc,
        [
            "Проверка датасета и YOLO-разметки.",
            "Обучение baseline YOLOv8n.",
            "Обучение YOLO11n и абляция входного разрешения 640/960.",
            "Обучение RT-DETR-L как альтернативного семейства.",
            "Экспорт лучшей модели в ONNX и сравнение скорости.",
            "Анализ ошибок и подготовка API-сервиса.",
        ],
    )

    doc.add_heading("Эксперименты и результаты", level=1)
    add_table(
        doc,
        ["id", "модель", "вход", "batch", "эпохи", "назначение"],
        [
            ["E1", "YOLOv8n", "416", "8", "3", "CPU quick baseline"],
            ["E2", "YOLO11n", "416", "8", "3", "современный легкий baseline"],
            ["E3", "YOLO11n", "960", "8", "40", "проверка малых объектов на GPU"],
            ["E4", "RT-DETR-L", "640", "8", "50", "архитектурная альтернатива на GPU"],
            ["E5", "YOLOv8n ONNX", "416", "1", "0", "деплой и скорость"],
        ],
        [0.55, 1.25, 0.7, 0.7, 0.7, 2.4],
    )
    doc.add_paragraph(
        "Из-за отсутствия CUDA/MPS в текущем окружении выполнены два быстрых CPU fine-tuning прогона. Также выполнен ONNX-экспорт выбранной модели. Полные GPU-конфиги сохранены в configs/."
    )
    add_table(
        doc,
        ["id", "модель", "mAP50", "mAP50-95", "precision", "recall", "latency"],
        [
            ["E1", "YOLOv8n quick", "0.12392", "0.06245", "0.33921", "0.17314", "51.0 ms"],
            ["E2", "YOLO11n quick", "0.11990", "0.06000", "0.41832", "0.17996", "55.7 ms"],
            ["E5", "YOLOv8n ONNX", "inherited", "inherited", "inherited", "inherited", "11.0 ms"],
        ],
        [0.45, 1.25, 0.75, 0.85, 0.8, 0.75, 0.85],
    )
    doc.add_paragraph(
        "Лучший текущий вариант по mAP50-95 и latency среди обученных PyTorch-моделей: YOLOv8n quick. Его веса сохранены в models/best.pt, ONNX-экспорт сохранен в models/best.onnx."
    )

    doc.add_heading("Метрики качества", level=1)
    add_table(
        doc,
        ["метрика", "смысл"],
        [
            ["mAP@0.5", "качество детекции при IoU 0.5"],
            ["mAP@0.5:0.95", "строгая интегральная метрика локализации"],
            ["precision", "доля корректных срабатываний среди предсказаний"],
            ["recall", "доля найденных объектов среди объектов разметки"],
            ["FP/FN", "ложные срабатывания и пропуски"],
            ["latency", "среднее время инференса"],
        ],
        [1.7, 4.5],
    )

    doc.add_heading("Деплой", level=1)
    doc.add_paragraph(
        "Реализован FastAPI-сервис с endpoint /health, /predict, /batch_predict, /stats и /metrics. История запусков сохраняется в SQLite: входной файл, backend, latency, число детекций и JSON-результат."
    )
    doc.add_paragraph(
        "Проект содержит Dockerfile и docker-compose.yml. Если models/best.pt отсутствует, сервис использует OpenCV fallback для smoke-test; после обучения автоматически загружает YOLO-веса."
    )

    doc.add_heading("Анализ ошибок", level=1)
    add_bullets(
        doc,
        [
            "Ложные срабатывания на тенях, заплатках асфальта и пятнах воды.",
            "Пропуск малых дефектов при низком разрешении.",
            "Смешение классов Pothole, medium-pothole и risk-pothole.",
            "Неполная локализация вытянутых трещин.",
            "Ошибки на частично перекрытых или размытых объектах.",
        ],
    )

    doc.add_heading("Вывод", level=1)
    doc.add_paragraph(
        "Проект оформлен как воспроизводимый практический пайплайн: подготовлены датасет, конфиги экспериментов, два реальных fine-tuning прогона, ONNX-экспорт, сервис инференса, Docker-запуск, история запросов, smoke-test и отчет. На CPU лучший результат показал YOLOv8n quick: mAP50-95=0.06245, latency около 51 ms/image в PyTorch и 11 ms/image в ONNX Runtime smoke-test. Для дальнейшего улучшения качества рекомендуется запустить полные конфиги на GPU и отдельно проверить объединение близких классов Pothole, medium-pothole и risk-pothole."
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Pothole Detection Practice")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
